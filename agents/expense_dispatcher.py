"""Expense Dispatcher — a small multi-agent workflow on top of ReceiptIQ.

    FetcherAgent  ->  AnalystAgent  ->  ComposerAgent  ->  DispatcherAgent
    (SQLite/API)      (categorise,       (formal expense     (Excel recipients,
                       anomalies,         summary: MD/HTML/    Gmail/SMTP or
                       aggregates)        CSV attachment)      --dry-run .eml)

Each agent is a plain Python class with a single responsibility and a `run(state)`
method that reads from / writes to a shared `WorkflowState`. The `Orchestrator`
executes them in order with per-agent retries and a run log, which is all the
"framework" this size of workflow needs. Swapping in LangChain/CrewAI is a matter
of wrapping each `run` as a tool/task (see README, Part 2).

Usage (from the repository root):
    python -m agents.make_sample_recipients                  # writes recipients.xlsx
    python -m agents.expense_dispatcher --recipients recipients.xlsx --dry-run
    python -m agents.expense_dispatcher --recipients recipients.xlsx   # real SMTP (see .env.example)
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import os
import re
import smtplib
import sqlite3
import sys
import time
from abc import ABC, abstractmethod
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from backend.app.config import settings  # noqa: E402  (re-uses .env loading + OpenRouter config)

log = logging.getLogger("expense_dispatcher")


# ============================================================================ state

@dataclass
class Recipient:
    name: str
    email: str
    department: str = ""


@dataclass
class WorkflowState:
    """Everything the agents share. Each agent only fills in its own section."""

    since: str | None = None
    until: str | None = None
    receipts: list[dict[str, Any]] = field(default_factory=list)      # FetcherAgent
    analysis: dict[str, Any] = field(default_factory=dict)            # AnalystAgent
    summary_markdown: str = ""                                        # ComposerAgent
    summary_html: str = ""
    summary_csv: str = ""
    subject: str = ""
    recipients: list[Recipient] = field(default_factory=list)         # DispatcherAgent
    dispatched: list[dict[str, Any]] = field(default_factory=list)
    run_log: list[str] = field(default_factory=list)

    def note(self, msg: str) -> None:
        stamp = datetime.now(timezone.utc).strftime("%H:%M:%S")
        self.run_log.append(f"[{stamp}] {msg}")
        log.info(msg)


class AgentError(Exception):
    pass


class Agent(ABC):
    name = "agent"
    max_attempts = 1

    @abstractmethod
    def run(self, state: WorkflowState) -> WorkflowState: ...


# ============================================================================ helpers

def _llm_json(system: str, user: str) -> dict | None:
    """Optional free-tier LLM call (OpenRouter). Returns parsed JSON or None on any failure."""
    if not settings.llm_enabled:
        return None
    import httpx

    try:
        r = httpx.post(
            f"{settings.openrouter_base_url}/chat/completions",
            headers={"Authorization": f"Bearer {settings.openrouter_api_key}", "X-Title": "ReceiptIQ agents"},
            json={
                "model": settings.openrouter_model,
                "temperature": 0,
                "response_format": {"type": "json_object"},
                "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
            },
            timeout=settings.openrouter_timeout_s,
        )
        r.raise_for_status()
        text = r.json()["choices"][0]["message"]["content"]
        return json.loads(text[text.find("{"): text.rfind("}") + 1])
    except Exception as exc:  # noqa: BLE001
        log.warning("LLM call skipped (%s); using deterministic path", exc)
        return None


CATEGORY_RULES: list[tuple[str, re.Pattern]] = [
    ("fuel", re.compile(r"\b(fuel|gas|petrol|diesel|pump|gallon|litre|liter|shell|bp|exxon|chevron|hp|indian oil|bharat)\b", re.I)),
    ("dining", re.compile(r"\b(bistro|cafe|caf[eé]|restaurant|grill|kitchen|diner|pizza|burger|bar|coffee|espresso|salad|salmon|tip|gratuity|server|table)\b", re.I)),
    ("travel", re.compile(r"\b(uber|ola|lyft|taxi|cab|airline|airways|flight|hotel|inn|lodge|rail|metro|parking|toll)\b", re.I)),
    ("groceries", re.compile(r"\b(grocery|supermarket|mart|market|foods|fresh|bigbasket|walmart|costco|kroger)\b", re.I)),
    ("retail", re.compile(r"\b(target|store|stores|shop|amazon|electronics|mouse|cable|notebook|office|supplies|stationery)\b", re.I)),
]


def categorise(receipt: dict[str, Any]) -> str:
    x = receipt.get("extracted") or {}
    blob = " ".join(
        [receipt.get("merchant_name") or "", " ".join(i.get("name", "") for i in x.get("line_items", [])), (x.get("raw_text") or "")[:400]]
    )
    for name, pattern in CATEGORY_RULES:
        if pattern.search(blob):
            return name
    return "other"


def _money(v: float | None) -> str:
    return "—" if v is None else f"${v:,.2f}"


# ============================================================================ agents

class FetcherAgent(Agent):
    """Loads extracted receipts from the ReceiptIQ SQLite DB (or the live REST API)."""

    name = "Fetcher"
    max_attempts = 3

    def __init__(self, db_path: Path | None = None, api_url: str | None = None):
        self.db_path = db_path
        self.api_url = api_url

    def _from_api(self) -> list[dict[str, Any]]:
        import httpx

        r = httpx.get(f"{self.api_url.rstrip('/')}/receipts", params={"limit": 1000}, timeout=30)
        r.raise_for_status()
        return r.json()["receipts"]

    def _from_db(self) -> list[dict[str, Any]]:
        if not self.db_path or not Path(self.db_path).exists():
            raise AgentError(f"Database not found: {self.db_path}")
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        try:
            rows = conn.execute("SELECT * FROM receipts ORDER BY date, id").fetchall()
        finally:
            conn.close()
        return [
            {
                "id": r["id"], "filename": r["filename"], "merchant_name": r["merchant_name"], "date": r["date"],
                "total_amount": r["total_amount"], "created_at": r["created_at"],
                "extracted": json.loads(r["extracted_json"]),
            }
            for r in rows
        ]

    def run(self, state: WorkflowState) -> WorkflowState:
        receipts = self._from_api() if self.api_url else self._from_db()
        if state.since:
            receipts = [r for r in receipts if (r.get("date") or "9999") >= state.since]
        if state.until:
            receipts = [r for r in receipts if (r.get("date") or "0000") <= state.until]
        state.receipts = receipts
        state.note(f"Fetcher: loaded {len(receipts)} receipt(s) from {'API ' + self.api_url if self.api_url else self.db_path}")
        return state


class AnalystAgent(Agent):
    """Categorises receipts, flags anomalies, and computes aggregates for the report."""

    name = "Analyst"
    max_attempts = 2

    def run(self, state: WorkflowState) -> WorkflowState:
        by_category: dict[str, float] = defaultdict(float)
        by_merchant: dict[str, float] = defaultdict(float)
        anomalies: list[dict[str, Any]] = []
        seen: Counter = Counter()
        total = tax = tip = 0.0

        for r in state.receipts:
            x = r.get("extracted") or {}
            r["category"] = categorise(r)
            amount = r.get("total_amount") or 0.0
            total += amount
            tax += x.get("tax") or 0.0
            tip += x.get("tip") or 0.0
            by_category[r["category"]] += amount
            by_merchant[r.get("merchant_name") or "Unknown"] += amount

            flags: list[str] = []
            if r.get("total_amount") is None:
                flags.append("missing total")
            if not r.get("date"):
                flags.append("missing date")
            if not r.get("merchant_name"):
                flags.append("missing merchant")
            if x.get("tip") and x.get("subtotal") and x["tip"] > 0.25 * x["subtotal"]:
                flags.append(f"tip is {x['tip'] / x['subtotal']:.0%} of subtotal")
            if amount and amount > 500:
                flags.append("high value (> $500), needs manager approval")
            flags.extend(w for w in x.get("warnings", []) if "does not match" in w or "sum to" in w)
            key = (r.get("merchant_name"), r.get("date"), r.get("total_amount"))
            seen[key] += 1
            if seen[key] > 1:
                flags.append("possible duplicate submission")
            if flags:
                anomalies.append({"id": r["id"], "merchant": r.get("merchant_name"), "flags": flags})
            r["flags"] = flags

        # Optional LLM pass: a one-paragraph narrative for the report. Deterministic fallback below.
        narrative = None
        if state.receipts:
            llm = _llm_json(
                "You are a finance analyst. Reply with JSON {\"narrative\": str} — one short formal paragraph "
                "summarising the expense report for a manager. No markdown.",
                json.dumps({"receipts": [{k: r.get(k) for k in ("merchant_name", "date", "total_amount", "category", "flags")} for r in state.receipts],
                            "totals": {"total": total, "tax": tax, "tip": tip}}),
            )
            narrative = (llm or {}).get("narrative")
        if not narrative:
            top = max(by_category.items(), key=lambda kv: kv[1])[0] if by_category else "n/a"
            narrative = (
                f"This report covers {len(state.receipts)} receipt(s) totalling {_money(total)}, of which "
                f"{_money(tax)} is tax and {_money(tip)} gratuity. The largest category is '{top}' at "
                f"{_money(by_category.get(top))}. {len(anomalies)} receipt(s) carry review flags."
            )

        state.analysis = {
            "count": len(state.receipts),
            "total": round(total, 2), "tax": round(tax, 2), "tip": round(tip, 2),
            "by_category": dict(sorted(by_category.items(), key=lambda kv: -kv[1])),
            "by_merchant": dict(sorted(by_merchant.items(), key=lambda kv: -kv[1])),
            "anomalies": anomalies,
            "narrative": narrative,
            "date_range": (
                min((r["date"] for r in state.receipts if r.get("date")), default=None),
                max((r["date"] for r in state.receipts if r.get("date")), default=None),
            ),
        }
        state.note(f"Analyst: total {_money(total)} across {len(by_category)} categories, {len(anomalies)} flagged")
        return state


class ComposerAgent(Agent):
    """Turns the analysis into a formal expense summary (Markdown + HTML + CSV attachment)."""

    name = "Composer"

    def run(self, state: WorkflowState) -> WorkflowState:
        a = state.analysis
        start, end = a.get("date_range", (None, None))
        period = f"{start} to {end}" if start and end else "all stored receipts"
        today = date.today().isoformat()
        state.subject = f"Expense Summary — {period} — {_money(a.get('total'))}"

        rows = [
            f"| {r['id']} | {r.get('date') or '—'} | {r.get('merchant_name') or '—'} | {r['category']} | "
            f"{_money((r.get('extracted') or {}).get('tax'))} | {_money((r.get('extracted') or {}).get('tip'))} | "
            f"{_money(r.get('total_amount'))} | {', '.join(r.get('flags', [])) or ''} |"
            for r in state.receipts
        ]
        cats = "\n".join(f"| {c} | {_money(v)} |" for c, v in a.get("by_category", {}).items())
        anomalies = "\n".join(f"- Receipt #{x['id']} ({x['merchant'] or 'unknown'}): {'; '.join(x['flags'])}" for x in a.get("anomalies", [])) or "- None"

        md = f"""# Expense Summary Report

**Period:** {period}
**Prepared:** {today} by ReceiptIQ Expense Dispatcher
**Receipts:** {a.get('count', 0)} · **Total:** {_money(a.get('total'))} · **Tax:** {_money(a.get('tax'))} · **Gratuity:** {_money(a.get('tip'))}

## Summary
{a.get('narrative', '')}

## Totals by category
| Category | Amount |
|---|---|
{cats or '| — | — |'}

## Receipts
| # | Date | Merchant | Category | Tax | Tip | Total | Flags |
|---|---|---|---|---|---|---|---|
{chr(10).join(rows) or '| — | — | — | — | — | — | — | — |'}

## Items requiring review
{anomalies}

---
*Generated automatically. Source data: ReceiptIQ SQLite store. Line-item detail is attached as CSV.*
"""
        state.summary_markdown = md
        state.summary_html = _markdown_to_html(md)

        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["receipt_id", "date", "merchant", "category", "item", "quantity", "price", "receipt_total", "flags"])
        for r in state.receipts:
            x = r.get("extracted") or {}
            items = x.get("line_items") or [{"name": "(no line items)", "quantity": None, "price": None}]
            for it in items:
                w.writerow([r["id"], r.get("date"), r.get("merchant_name"), r["category"], it.get("name"), it.get("quantity"), it.get("price"), r.get("total_amount"), "; ".join(r.get("flags", []))])
        state.summary_csv = buf.getvalue()
        state.note(f"Composer: report '{state.subject}' ({len(md)} chars, CSV {len(state.summary_csv)} bytes)")
        return state


def _markdown_to_html(md: str) -> str:
    """Minimal Markdown -> HTML for headings, bold, bullet lists and pipe tables (no dependency)."""
    html: list[str] = ["<html><body style='font-family:system-ui,sans-serif;max-width:860px'>"]
    table: list[str] = []

    def flush_table() -> None:
        if table:
            head, *body = [row.strip("|").split("|") for row in table if not re.match(r"^\|?\s*-", row)]
            html.append("<table border='1' cellpadding='6' style='border-collapse:collapse'>")
            html.append("<tr>" + "".join(f"<th>{c.strip()}</th>" for c in head) + "</tr>")
            for r in body:
                html.append("<tr>" + "".join(f"<td>{c.strip()}</td>" for c in r) + "</tr>")
            html.append("</table>")
            table.clear()

    for line in md.splitlines():
        if line.startswith("|"):
            table.append(line)
            continue
        flush_table()
        line = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line)
        line = re.sub(r"\*(.+?)\*", r"<i>\1</i>", line)
        if line.startswith("# "):
            html.append(f"<h1>{line[2:]}</h1>")
        elif line.startswith("## "):
            html.append(f"<h2>{line[3:]}</h2>")
        elif line.startswith("- "):
            html.append(f"<li>{line[2:]}</li>")
        elif line.strip() == "---":
            html.append("<hr>")
        elif line.strip():
            html.append(f"<p>{line.rstrip()}</p>")
    flush_table()
    html.append("</body></html>")
    return "\n".join(html)


class DispatcherAgent(Agent):
    """Reads recipients from an Excel sheet and emails the report (SMTP) or writes .eml files (dry run)."""

    name = "Dispatcher"
    max_attempts = 3

    def __init__(self, recipients_xlsx: Path, dry_run: bool, outbox: Path):
        self.recipients_xlsx = recipients_xlsx
        self.dry_run = dry_run
        self.outbox = outbox

    def load_recipients(self) -> list[Recipient]:
        from openpyxl import load_workbook

        wb = load_workbook(self.recipients_xlsx, read_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().lower() for h in next(rows)]
        col = {name: header.index(name) for name in ("name", "email", "department") if name in header}
        if "email" not in col:
            raise AgentError("recipients sheet needs an 'email' column (optional: name, department)")
        recipients: list[Recipient] = []
        for row in rows:
            email = str(row[col["email"]] or "").strip() if row else ""
            if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
                continue
            recipients.append(Recipient(
                name=str(row[col["name"]] or "").strip() if "name" in col else email,
                email=email,
                department=str(row[col["department"]] or "").strip() if "department" in col else "",
            ))
        return recipients

    def build_message(self, state: WorkflowState, rcpt: Recipient, sender: str) -> EmailMessage:
        msg = EmailMessage()
        msg["Subject"] = state.subject
        msg["From"] = sender
        msg["To"] = f"{rcpt.name} <{rcpt.email}>" if rcpt.name else rcpt.email
        greeting = f"Dear {rcpt.name or 'colleague'},\n\nPlease find the expense summary below.\n\n"
        msg.set_content(greeting + state.summary_markdown)
        msg.add_alternative(state.summary_html.replace("<body", f"<body><p>Dear {rcpt.name or 'colleague'},</p>", 1), subtype="html")
        msg.add_attachment(state.summary_csv.encode(), maintype="text", subtype="csv", filename="expense_line_items.csv")
        return msg

    def run(self, state: WorkflowState) -> WorkflowState:
        state.recipients = self.load_recipients()
        if not state.recipients:
            raise AgentError("No valid recipients found in the Excel sheet")
        sender = os.getenv("SMTP_FROM") or os.getenv("SMTP_USER") or "receiptiq@example.com"

        if self.dry_run:
            self.outbox.mkdir(parents=True, exist_ok=True)
            for rcpt in state.recipients:
                msg = self.build_message(state, rcpt, sender)
                path = self.outbox / f"{re.sub(r'[^a-z0-9]+', '_', rcpt.email.lower())}.eml"
                path.write_bytes(bytes(msg))
                state.dispatched.append({"to": rcpt.email, "status": "dry-run", "file": str(path)})
            state.note(f"Dispatcher: DRY RUN — wrote {len(state.recipients)} .eml file(s) to {self.outbox}")
            return state

        host, port = os.getenv("SMTP_HOST", "smtp.gmail.com"), int(os.getenv("SMTP_PORT", "587"))
        user, password = os.getenv("SMTP_USER"), os.getenv("SMTP_PASSWORD")
        if not (user and password):
            raise AgentError("SMTP_USER and SMTP_PASSWORD are required unless --dry-run is used")
        smtp_cls = smtplib.SMTP_SSL if port == 465 else smtplib.SMTP
        with smtp_cls(host, port, timeout=30) as smtp:
            if port != 465:
                smtp.starttls()
            smtp.login(user, password)
            for rcpt in state.recipients:
                try:
                    smtp.send_message(self.build_message(state, rcpt, sender))
                    state.dispatched.append({"to": rcpt.email, "status": "sent"})
                except smtplib.SMTPException as exc:
                    state.dispatched.append({"to": rcpt.email, "status": f"failed: {exc}"})
        sent = sum(d["status"] == "sent" for d in state.dispatched)
        state.note(f"Dispatcher: sent {sent}/{len(state.recipients)} email(s) via {host}:{port}")
        return state


# ============================================================================ orchestrator

class Orchestrator:
    """Runs agents sequentially with per-agent retry and a shared state. Fails fast on a dead agent."""

    def __init__(self, agents: list[Agent]):
        self.agents = agents

    def run(self, state: WorkflowState) -> WorkflowState:
        for agent in self.agents:
            for attempt in range(1, agent.max_attempts + 1):
                try:
                    t0 = time.perf_counter()
                    state = agent.run(state)
                    state.note(f"{agent.name}: done in {time.perf_counter() - t0:.2f}s")
                    break
                except AgentError:
                    raise
                except Exception as exc:  # noqa: BLE001
                    state.note(f"{agent.name}: attempt {attempt}/{agent.max_attempts} failed: {exc}")
                    if attempt == agent.max_attempts:
                        raise AgentError(f"{agent.name} failed after {attempt} attempt(s): {exc}") from exc
                    time.sleep(0.5 * attempt)
        return state


# ============================================================================ CLI

def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="Compose and dispatch an expense summary from ReceiptIQ receipts.")
    p.add_argument("--recipients", required=True, type=Path, help="Excel sheet with columns: name, email, department")
    p.add_argument("--db", type=Path, default=settings.db_path, help="ReceiptIQ SQLite DB (default: data/receipts.db)")
    p.add_argument("--api-url", help="Fetch from a running ReceiptIQ API instead of the DB, e.g. http://localhost:8000")
    p.add_argument("--since", help="Only receipts dated >= YYYY-MM-DD")
    p.add_argument("--until", help="Only receipts dated <= YYYY-MM-DD")
    p.add_argument("--dry-run", action="store_true", help="Write .eml files to --outbox instead of sending")
    p.add_argument("--outbox", type=Path, default=PROJECT_ROOT / "outbox")
    p.add_argument("--print-report", action="store_true", help="Print the Markdown report to stdout")
    args = p.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")

    state = WorkflowState(since=args.since, until=args.until)
    orchestrator = Orchestrator([
        FetcherAgent(db_path=args.db, api_url=args.api_url),
        AnalystAgent(),
        ComposerAgent(),
        DispatcherAgent(args.recipients, dry_run=args.dry_run, outbox=args.outbox),
    ])
    try:
        state = orchestrator.run(state)
    except AgentError as exc:
        log.error("workflow aborted: %s", exc)
        return 1

    if args.print_report:
        print(state.summary_markdown)
    print(json.dumps({"subject": state.subject, "receipts": state.analysis.get("count"), "total": state.analysis.get("total"),
                      "dispatched": state.dispatched}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
